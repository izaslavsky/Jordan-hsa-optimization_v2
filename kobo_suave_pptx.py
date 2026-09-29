"""
Build PPTX for the Kobo-SuAVE workshop session.
Usage: python kobo_suave_pptx.py [<output.pptx>]
"""
import re, sys, io, zipfile
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

_REPO    = Path(__file__).parent
GRAPHICS = _REPO / "manuscript" / "kobo_suave_graphics"

# ── Image map: slide title → (image path, layout hint) ──────────────────────
# layout hints: "wide" = full content width; "right" = right half; "full" = full slide

IMAGE_MAP: dict[str, tuple[Path, str]] = {
    # ── Block 1: Introduction ────────────────────────────────────────────────
    "The Challenge with Field Survey Data":  (GRAPHICS / "ks_health_pathway.png",      "right"),
    "The GC3WEFH Pipeline: Three Components":(GRAPHICS / "ks_pipeline.png",            "wide"),
    "What We Are Exploring Today":           (GRAPHICS / "ks_health_pathway.png",      "right"),
    "Session Roadmap":                       (GRAPHICS / "ks_session_roadmap.png",     "wide"),
    "Open Your Browser Now":                 (GRAPHICS / "ks_pipeline.png",            "right"),

    # ── Block 2: Exploring in SuAVE ─────────────────────────────────────────
    "The SuAVE Interface at a Glance":       (GRAPHICS / "ks_suave_interface.png",     "wide"),
    "What the Icons Mean":                   (GRAPHICS / "ks_health_pathway.png",      "right"),
    "Key Variables for Exploration":         (GRAPHICS / "ks_health_pathway.png",      "right"),
    "The Views We Will Use Today":           (GRAPHICS / "ks_views_overview.png",      "wide"),
    "The Core Exploratory Question":         (GRAPHICS / "ks_health_pathway.png",      "wide"),
    "Sharing and Downloading":               (GRAPHICS / "ks_suave_interface.png",     "right"),

    # ── Block 3: Authoring ───────────────────────────────────────────────────
    "The Five-Minute Publishing Workflow":   (GRAPHICS / "ks_pipeline.png",            "right"),
    "Column Annotation Reference":           (GRAPHICS / "ks_column_annotations.png",  "wide"),
    "Configuring Views and Icons":           (GRAPHICS / "ks_suave_interface.png",     "right"),
    "Cloning a Public Survey":               (GRAPHICS / "ks_suave_interface.png",     "right"),
    "Access Control and Annotation":         (GRAPHICS / "ks_suave_interface.png",     "right"),

    # ── Block 4: KoboToolbox ─────────────────────────────────────────────────
    "KoboToolbox Overview":                  (GRAPHICS / "ks_kobo_form.png",           "wide"),
    "Designing a Climate-Health Survey: What to Include":
                                             (GRAPHICS / "ks_kobo_form.png",           "right"),
    "Live Demo: Build a Five-Question Form": (GRAPHICS / "ks_kobo_form.png",           "right"),
    "Deploying and Collecting":              (GRAPHICS / "ks_kobo_form.png",           "right"),

    # ── Block 5: Bridge ──────────────────────────────────────────────────────
    "How the Bridge Works":                  (GRAPHICS / "ks_bridge_mechanism.png",    "wide"),
    "Configuration Steps":                   (GRAPHICS / "ks_bridge_mechanism.png",    "right"),
    "Field Mapping: Kobo Question Types to SuAVE Columns":
                                             (GRAPHICS / "ks_kobo_form.png",           "right"),
    "Live Test: Submit and Watch It Appear": (GRAPHICS / "ks_bridge_mechanism.png",    "right"),

    # ── Block 6: Jupyter / Streamlit ─────────────────────────────────────────
    "When the Built-In Views Are Not Enough":(GRAPHICS / "ks_jupyter_flow.png",        "right"),
    "Jupyter Integration":                   (GRAPHICS / "ks_jupyter_flow.png",        "wide"),
    "Streamlit Integration":                 (GRAPHICS / "ks_jupyter_flow.png",        "right"),
    "Pre-Built Templates and Custom Extensions":
                                             (GRAPHICS / "ks_views_overview.png",      "right"),

    # ── Block 7: Q&A ─────────────────────────────────────────────────────────
    "Explore, Share, and Ask":               (GRAPHICS / "ks_suave_interface.png",     "right"),
}

# ── Colour palette ───────────────────────────────────────────────────────────
C_TITLE      = RGBColor(0x1A, 0x37, 0x5E)
C_ACCENT     = RGBColor(0x2E, 0x86, 0xC1)
C_BODY       = RGBColor(0x1C, 0x1C, 0x1C)
C_MUTED      = RGBColor(0x55, 0x55, 0x55)
C_TABLE_HDR  = RGBColor(0x2E, 0x86, 0xC1)
C_TABLE_ROW1 = RGBColor(0xEB, 0xF5, 0xFB)
C_TABLE_ROW2 = RGBColor(0xF8, 0xF9, 0xF9)
C_WHITE      = RGBColor(0xFF, 0xFF, 0xFF)

SLIDE_W = Emu(12192000)
SLIDE_H = Emu(6858000)

FS_TITLE = Pt(28)
FS_BODY  = Pt(16)
FS_SMALL = Pt(13)
FS_TABLE = Pt(12)

M        = Inches(0.55)
TOP_BAND = Inches(1.15)


# ─────────────────────────────────────────────────────────────────────────────
# PARSING  (same pattern as md_to_pptx.py)
# ─────────────────────────────────────────────────────────────────────────────

def parse_slides(md_text: str) -> list[dict]:
    slides = []
    slide_re   = re.compile(r'^### Slide\s+[\d.]+\s+[—–-]+\s+(.+)$', re.MULTILINE)
    session_re = re.compile(r'^# (SESSION \d+.*)',  re.MULTILINE)
    section_re = re.compile(r'^## (PART \d+.*?)(?:\s*[—–-].*)?$', re.MULTILINE)

    events = []
    for m in session_re.finditer(md_text):
        events.append(('session', m.start(), m.group(1).strip()))
    for m in section_re.finditer(md_text):
        events.append(('section', m.start(), m.group(1).strip()))
    for m in slide_re.finditer(md_text):
        events.append(('slide', m.start(), m.group(1).strip(), m.end()))
    events.sort(key=lambda e: e[1])

    slide_starts = [(i, e) for i, e in enumerate(events) if e[0] == 'slide']

    for idx, (ev_idx, event) in enumerate(slide_starts):
        title      = event[2]
        body_start = event[3]
        body_end   = slide_starts[idx + 1][1][1] if idx + 1 < len(slide_starts) else len(md_text)
        body       = md_text[body_start:body_end]

        wbn = sec = ""
        for e in events[:ev_idx + 1]:
            if e[0] == 'session': wbn = e[2]
            elif e[0] == 'section': sec = e[2]

        slide = _parse_body(title, body)
        slide['session'] = wbn
        slide['section'] = sec
        slides.append(slide)

    return slides


def _parse_body(title: str, body: str) -> dict:
    lines, layout, visual = body.split('\n'), "", ""
    notes_lines, content_lines = [], []
    in_notes = False

    for line in lines:
        s = line.rstrip()
        if s in ('---', '----'): continue
        if not layout and s.lower().startswith('layout:'):
            layout = s[len('layout:'):].strip(); continue
        if not visual and s.lower().startswith('visual:'):
            visual = s[len('visual:'):].strip(); continue
        if s.lower().startswith('notes:'):
            in_notes = True
            rest = s[len('notes:'):].strip()
            if rest: notes_lines.append(rest)
            continue
        (notes_lines if in_notes else content_lines).append(s)

    while content_lines and not content_lines[-1].strip():
        content_lines.pop()

    return {
        'title': title, 'layout': layout, 'visual': visual,
        'content_lines': content_lines, 'notes': ' '.join(notes_lines).strip(),
    }


# ─────────────────────────────────────────────────────────────────────────────
# CONTENT HELPERS
# ─────────────────────────────────────────────────────────────────────────────

def clean_inline(text: str) -> str:
    text = re.sub(r'\*\*(.+?)\*\*', r'\1', text)
    text = re.sub(r'`(.+?)`', r'\1', text)
    text = re.sub(r'\[(.+?)\]\(.+?\)', r'\1', text)
    return text

def is_table_row(line: str) -> bool:
    return line.strip().startswith('|') and '|' in line[1:]

def is_separator_row(line: str) -> bool:
    return bool(re.match(r'^\|[-| :]+\|$', line.strip()))

def parse_table(content_lines):
    tbl = [l for l in content_lines if is_table_row(l) and not is_separator_row(l)]
    if not tbl: return [], []
    def split_row(l): return [c.strip() for c in l.strip().strip('|').split('|')]
    return split_row(tbl[0]), [split_row(l) for l in tbl[1:]]

def is_bullet(line: str) -> bool:
    return line.strip().startswith('•') or line.strip().startswith('-')

def extract_bullets(content_lines):
    bullets = []
    for line in content_lines:
        s = line.strip()
        if is_bullet(s):
            bullets.append(clean_inline(s.lstrip('•').lstrip('-').strip()))
        elif s and not is_table_row(s):
            bullets.append(clean_inline(s))
    return [b for b in bullets if b]


# ─────────────────────────────────────────────────────────────────────────────
# PPTX BUILDING BLOCKS
# ─────────────────────────────────────────────────────────────────────────────

def blank_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])

def add_title_band(slide, title: str, section: str = "", session: str = "") -> None:
    bg = slide.shapes.add_shape(1, Emu(0), Emu(0), SLIDE_W, TOP_BAND)
    bg.fill.solid(); bg.fill.fore_color.rgb = C_TITLE; bg.line.fill.background()

    txb = slide.shapes.add_textbox(M, Inches(0.12), SLIDE_W - 2 * M, TOP_BAND - Inches(0.12))
    tf = txb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.LEFT
    run = p.add_run()
    run.text = clean_inline(title); run.font.size = FS_TITLE
    run.font.bold = True; run.font.color.rgb = C_WHITE

    if section or session:
        label = f"{session}  ·  {section}" if (session and section) else (session or section)
        ltxb = slide.shapes.add_textbox(M, Inches(0.83), SLIDE_W - 2 * M, Inches(0.30))
        lp = ltxb.text_frame.paragraphs[0]; lp.alignment = PP_ALIGN.RIGHT
        lr = lp.add_run(); lr.text = label; lr.font.size = Pt(9)
        lr.font.color.rgb = RGBColor(0xA9, 0xCC, 0xE3)

def add_visual_cue(slide, visual: str) -> None:
    if not visual or visual.lower() in ('none', 'none — clean text only'):
        return
    txb = slide.shapes.add_textbox(M, SLIDE_H - Inches(0.38), SLIDE_W - 2 * M, Inches(0.34))
    p = txb.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.LEFT
    run = p.add_run()
    run.text = f"[Visual: {visual[:160]}{'…' if len(visual) > 160 else ''}]"
    run.font.size = Pt(8); run.font.italic = True; run.font.color.rgb = C_MUTED

def add_image_shape(slide, img_path: Path, left, top, max_width, max_height) -> None:
    from PIL import Image as _PIL
    try:
        with _PIL.open(str(img_path)) as im:
            iw, ih = im.size
    except Exception:
        slide.shapes.add_picture(str(img_path), left, top, max_width)
        return
    asp_img = iw / ih; asp_box = max_width / max_height
    if asp_img >= asp_box:
        fw = int(max_width); fh = int(max_width / asp_img)
    else:
        fh = int(max_height); fw = int(max_height * asp_img)
    al = int(left + (max_width  - fw) // 2)
    at = int(top  + (max_height - fh) // 2)
    slide.shapes.add_picture(str(img_path), al, at, fw, fh)

def add_bullets_box(slide, bullets, top=None, height=None, font_size=None,
                    left=None, width=None) -> None:
    t = top   if top   is not None else (TOP_BAND + Inches(0.2))
    h = height if height is not None else (SLIDE_H - t - Inches(0.45))
    fs = font_size or FS_BODY
    l = left  if left  is not None else M
    w = width if width is not None else (SLIDE_W - 2 * M)
    txb = slide.shapes.add_textbox(l, t, w, h)
    tf = txb.text_frame; tf.word_wrap = True
    first = True
    for bullet in bullets:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_before = Pt(4)
        is_sub = bullet.startswith('  ')
        bullet = bullet.strip()
        p.level = 1 if is_sub else 0
        run = p.add_run()
        run.text = f"• {bullet}" if not is_sub else f"  – {bullet}"
        run.font.size = fs if not is_sub else Pt(fs.pt - 2)
        run.font.color.rgb = C_BODY

def add_table_shape(slide, headers, rows, top=None, height=None) -> None:
    t = top or (TOP_BAND + Inches(0.25))
    h = height if height is not None else (SLIDE_H - t - Inches(0.45))
    ncols = len(headers); col_w = (SLIDE_W - 2 * M) / ncols
    tbl = slide.shapes.add_table(len(rows) + 1, ncols, M, t,
                                  SLIDE_W - 2 * M, h).table
    for ci in range(ncols):
        tbl.columns[ci].width = int(col_w)

    def set_cell(row, col, text, bold=False, bg=None, fs=None):
        cell = tbl.cell(row, col)
        cell.text = clean_inline(text)
        for para in cell.text_frame.paragraphs:
            for run in para.runs:
                run.font.size = fs or FS_TABLE
                run.font.bold = bold
                run.font.color.rgb = C_WHITE if bold else C_BODY
        if bg:
            cell.fill.solid(); cell.fill.fore_color.rgb = bg

    for ci, hdr in enumerate(headers):
        set_cell(0, ci, hdr, bold=True, bg=C_TABLE_HDR)
    for ri, row in enumerate(rows):
        bg = C_TABLE_ROW1 if ri % 2 == 0 else C_TABLE_ROW2
        for ci, cell_text in enumerate(row):
            set_cell(ri + 1, ci, cell_text, bg=bg)


# ─────────────────────────────────────────────────────────────────────────────
# SECTION DIVIDER
# ─────────────────────────────────────────────────────────────────────────────

def add_section_divider(prs, section_title: str) -> None:
    slide = blank_slide(prs)
    bg = slide.shapes.add_shape(1, Emu(0), Emu(0), SLIDE_W, SLIDE_H)
    bg.fill.solid(); bg.fill.fore_color.rgb = C_TITLE; bg.line.fill.background()

    txb = slide.shapes.add_textbox(
        Inches(1.0), Inches(2.3), SLIDE_W - Inches(2.0), Inches(2.0)
    )
    tf = txb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = section_title
    run.font.size = Pt(34); run.font.bold = True
    run.font.color.rgb = C_WHITE

    accent = slide.shapes.add_shape(
        1, Inches(4.5), Inches(4.5), Inches(4.2), Inches(0.06)
    )
    accent.fill.solid(); accent.fill.fore_color.rgb = C_ACCENT
    accent.line.fill.background()


# ─────────────────────────────────────────────────────────────────────────────
# SLIDE RENDERER
# ─────────────────────────────────────────────────────────────────────────────

CONTENT_TOP = TOP_BAND + Inches(0.18)
CONTENT_H   = SLIDE_H  - CONTENT_TOP - Inches(0.45)
CONTENT_W   = SLIDE_W  - 2 * M


def render_slide(prs, slide_data: dict) -> None:
    title   = slide_data['title']
    layout  = slide_data.get('layout', '').lower()
    visual  = slide_data.get('visual', '')
    content = slide_data.get('content_lines', [])
    notes   = slide_data.get('notes', '')
    section = slide_data.get('section', '')
    session = slide_data.get('session', '')

    img_entry = IMAGE_MAP.get(title)
    img_path  = img_entry[0] if img_entry else None
    img_hint  = img_entry[1] if img_entry else None

    if img_path and not img_path.exists():
        img_path = None

    headers, rows   = parse_table(content)
    has_table        = bool(headers)
    bullets_all      = extract_bullets(content)

    slide = blank_slide(prs)
    add_title_band(slide, title, section=section, session=session)

    if img_path and img_hint == "full":
        add_image_shape(slide, img_path, Emu(0), Emu(0), SLIDE_W, SLIDE_H)
        add_title_band(slide, title, section=section, session=session)
        if notes:
            _add_notes(slide, notes)
        return

    if img_path and img_hint == "wide":
        IMG_H = int(CONTENT_H * 0.52) if (has_table or bullets_all) else int(CONTENT_H * 0.85)
        add_image_shape(slide, img_path, M, CONTENT_TOP, CONTENT_W, IMG_H)
        BELOW_TOP = int(CONTENT_TOP + IMG_H + Inches(0.08))
        BELOW_H   = int(SLIDE_H - BELOW_TOP - Inches(0.45))
        if has_table:
            add_table_shape(slide, headers, rows, top=BELOW_TOP, height=BELOW_H)
        elif bullets_all:
            n = len(bullets_all)
            fs = Pt(22) if n <= 3 else (Pt(19) if n <= 5 else FS_BODY)
            add_bullets_box(slide, bullets_all, top=BELOW_TOP, height=BELOW_H, font_size=fs)

    elif img_path and img_hint == "right":
        IMG_W = int(CONTENT_W * 0.47)
        TXT_W = int(CONTENT_W * 0.48)
        add_image_shape(slide, img_path,
                        int(SLIDE_W - M - IMG_W), CONTENT_TOP,
                        IMG_W, CONTENT_H)
        if has_table:
            tbl_top = int(CONTENT_TOP + Inches(0.05))
            tbl_h   = int(CONTENT_H - Inches(0.1))
            slide.shapes.add_table(
                len(rows) + 1, len(headers),
                M, tbl_top, TXT_W, tbl_h
            )
            add_table_shape(slide, headers, rows, top=tbl_top, height=tbl_h)
        elif bullets_all:
            n = len(bullets_all)
            fs = Pt(19) if n <= 3 else (Pt(16) if n <= 5 else Pt(14))
            add_bullets_box(slide, bullets_all, top=CONTENT_TOP, height=CONTENT_H,
                            left=M, width=TXT_W, font_size=fs)

    else:
        if has_table:
            tbl_top = int(CONTENT_TOP + Inches(0.05))
            tbl_h   = int(SLIDE_H - tbl_top - Inches(0.45))
            if bullets_all:
                add_bullets_box(slide, bullets_all, top=CONTENT_TOP,
                                height=int(Inches(0.8)), font_size=Pt(13))
                tbl_top = int(CONTENT_TOP + Inches(0.9))
                tbl_h   = int(SLIDE_H - tbl_top - Inches(0.45))
            add_table_shape(slide, headers, rows, top=tbl_top, height=tbl_h)
        elif bullets_all:
            n = len(bullets_all)
            fs = Pt(22) if n <= 3 else (Pt(19) if n <= 5 else FS_BODY)
            add_bullets_box(slide, bullets_all, font_size=fs)

    if not has_table and not bullets_all and not img_path:
        add_visual_cue(slide, visual)

    if notes:
        _add_notes(slide, notes)


def _add_notes(slide, notes: str) -> None:
    notes_slide = slide.notes_slide
    tf = notes_slide.notes_text_frame
    tf.text = notes


# ─────────────────────────────────────────────────────────────────────────────
# NOTES-MASTER FIX
# ─────────────────────────────────────────────────────────────────────────────

def _fix_notes_master(pptx_path: Path) -> None:
    with zipfile.ZipFile(pptx_path, 'r') as zin:
        rels_raw = zin.read('ppt/_rels/presentation.xml.rels')
        import xml.etree.ElementTree as _ET
        rels = _ET.fromstring(rels_raw)
        notes_rId = None
        for rel in rels:
            if 'notesMaster' in rel.get('Type', ''):
                notes_rId = rel.get('Id'); break
        if not notes_rId:
            return
        prs_raw = zin.read('ppt/presentation.xml').decode('utf-8')
        if 'notesMasterIdLst' in prs_raw:
            return
        insert_xml = (
            f'<p:notesMasterIdLst xmlns:p="http://schemas.openxmlformats.org/'
            f'presentationml/2006/main" xmlns:r="http://schemas.openxmlformats.org/'
            f'officeDocument/2006/relationships">'
            f'<p:notesMasterId r:id="{notes_rId}"/></p:notesMasterIdLst>'
        )
        prs_fixed = prs_raw.replace(
            '</p:sldMasterIdLst><p:sldIdLst>',
            f'</p:sldMasterIdLst>{insert_xml}<p:sldIdLst>'
        )
        if prs_fixed == prs_raw:
            return
    buf = io.BytesIO()
    with zipfile.ZipFile(pptx_path, 'r') as zin2, \
         zipfile.ZipFile(buf, 'w', compression=zipfile.ZIP_DEFLATED) as zout:
        for item in zin2.infolist():
            data = prs_fixed.encode('utf-8') if item.filename == 'ppt/presentation.xml' \
                   else zin2.read(item.filename)
            zout.writestr(item, data)
    pptx_path.write_bytes(buf.getvalue())


# ─────────────────────────────────────────────────────────────────────────────
# MAIN
# ─────────────────────────────────────────────────────────────────────────────

def build(md_path: Path, out_path: Path) -> None:
    md_text = md_path.read_text(encoding='utf-8')
    slides  = parse_slides(md_text)

    prs = Presentation()
    prs.slide_width  = SLIDE_W
    prs.slide_height = SLIDE_H

    seen_sections = set()
    for slide_data in slides:
        sec = slide_data.get('section', '')
        if sec and sec not in seen_sections:
            add_section_divider(prs, sec)
            seen_sections.add(sec)
        render_slide(prs, slide_data)

    prs.save(str(out_path))
    _fix_notes_master(out_path)
    print(f"Saved {len(slides)} slides → {out_path}")


if __name__ == '__main__':
    md_in  = _REPO / "manuscript" / "kobo_suave_slides.md"
    pptx_out = Path(sys.argv[1]) if len(sys.argv) > 1 \
               else _REPO / "manuscript" / "kobo_suave_v1.pptx"
    build(md_in, pptx_out)
